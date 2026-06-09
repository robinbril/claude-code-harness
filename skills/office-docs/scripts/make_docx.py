"""Starter: een net Word-document met kop, tekst en een gearceerde tabel.
uv run --with python-docx python make_docx.py"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY = RGBColor(0x0E, 0x1B, 0x26)
TEAL_LT = "E3F6F7"


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)

p = doc.add_paragraph()
r = p.add_run("Titel van het document"); r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY

doc.add_paragraph("Een inleidende alinea met gewone tekst. Houd het kort en helder.")

t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
for label, value in [("Onderwerp", "Waarde"), ("Tweede", "Nog een waarde")]:
    cells = t.add_row().cells
    shade(cells[0], TEAL_LT)
    cells[0].paragraphs[0].add_run(label).font.bold = True
    cells[1].paragraphs[0].add_run(value)

out = Path.home() / "Downloads" / "voorbeeld.docx"
try:
    doc.save(out)
except PermissionError:
    out = out.with_name(out.stem + " (nieuw)" + out.suffix); doc.save(out)
print("OK ->", out)
