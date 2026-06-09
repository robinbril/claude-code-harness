"""Starter: een eenvoudige PowerPoint. uv run --with python-pptx python make_pptx.py
Voor een vaste huisstijl: maak een eigen wrapper met je palet en fonts."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x22, 0x22, 0x22)
ACCENT = RGBColor(0x2D, 0x6C, 0xDF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])


def shadow_off(sh):
    if sh._element.spPr.find(qn("a:effectLst")) is None:
        sh._element.spPr.append(sh._element.spPr.makeelement(qn("a:effectLst"), {}))


def textbox(x, y, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = tb.text_frame.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Segoe UI"
    return tb


bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.18))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); shadow_off(bar)
textbox(0.8, 2.6, 11.5, 1.2, "Titel van de presentatie", size=44, bold=True)
textbox(0.85, 3.9, 11.5, 0.8, "Ondertitel of korte toelichting", size=20, color=RGBColor(0x66, 0x66, 0x66))

out = Path.home() / "Downloads" / "voorbeeld.pptx"
try:
    prs.save(out)
except PermissionError:
    out = out.with_name(out.stem + " (nieuw)" + out.suffix); prs.save(out)
print("OK ->", out)
