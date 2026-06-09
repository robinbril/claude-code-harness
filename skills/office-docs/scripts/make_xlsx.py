"""Starter: een nette Excel met kop, kleuren, dropdown en wrap. uv run --with openpyxl python make_xlsx.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

NAVY, GREEN, AMBER, RED = "0E1B26", "C6EFCE", "FFE9A8", "FFC7CE"
thin = Side(style="thin", color="D9E1E3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = Workbook(); ws = wb.active; ws.title = "Overzicht"; ws.sheet_view.showGridLines = False
headers = ["#", "Taak", "Status"]
widths = [5, 60, 18]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(1, i, h); c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(bold=True, color="FFFFFF"); c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

rows = [("Eerste taak met wat langere omschrijving die netjes moet wrappen", "Te doen"),
        ("Tweede taak", "Verwerkt")]
for idx, (taak, status) in enumerate(rows, 1):
    r = idx + 1
    ws.cell(r, 1, idx).alignment = Alignment(horizontal="center")
    ws.cell(r, 2, taak).alignment = WRAP
    ws.cell(r, 3, status).alignment = Alignment(horizontal="center")
    for col in range(1, 4):
        ws.cell(r, col).border = BORDER
    # rijhoogte zelf zetten voor wraptekst (openpyxl groeit niet mee)
    lines = max(1, -(-len(taak) // 58))
    ws.row_dimensions[r].height = 15 * lines + 4

last = len(rows) + 1
dv = DataValidation(type="list", formula1='"Te doen,Bezig,Verwerkt"', allow_blank=True)
dv.add(f"C2:C{last}"); ws.add_data_validation(dv)
for val, fill in (("Verwerkt", GREEN), ("Bezig", AMBER), ("Te doen", RED)):
    ws.conditional_formatting.add(f"C2:C{last}",
        CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=fill)))
ws.auto_filter.ref = f"A1:C{last}"

out = Path.home() / "Downloads" / "voorbeeld.xlsx"
try:
    wb.save(out)
except PermissionError:
    out = out.with_name(out.stem + " (nieuw)" + out.suffix); wb.save(out)
print("OK ->", out)
